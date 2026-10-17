#!/usr/bin/env python3
"""Build four editable 8-slide Chinese paper-report decks in two batches.

Phase 1 creates slides 1-5. Phase 2 loads the phase-1 deck and appends slides 6-8.
Text remains editable; paper figures and math formulas are PNG elements.
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

import fitz
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path("/mnt/data/zhangshuai/LLMs")
BASE = ROOT / "PPT/李老师组会/Agent动作投机Survey"
ASSET = BASE / "assets"
DECK = BASE / "decks"
LOGO = ASSET / "whu-cse-logo.png"
DSP = BASE / "sources/DSP/latex"
SPORK = BASE / "sources/SPORK/latex"
JOINT = BASE / "sources/Joint-RL/latex"
PASTE = ROOT / "paper/Agent系列latex论文/Agent动作投机/Parallelizing Tool Execution and LLM Generation for Low-Latency Agent Serving"

FONT_CN = "Microsoft YaHei"
FONT_EQ = "Times New Roman"
RED = RGBColor(166, 0, 0)
RED2 = RGBColor(205, 45, 45)
PALE_RED = RGBColor(255, 245, 245)
TEAL = RGBColor(0, 128, 123)
TEAL2 = RGBColor(30, 160, 151)
PALE_TEAL = RGBColor(241, 251, 249)
BLUE = RGBColor(36, 94, 168)
PALE_BLUE = RGBColor(244, 248, 255)
DARK = RGBColor(26, 34, 48)
MID = RGBColor(78, 87, 102)
LIGHT = RGBColor(224, 228, 234)
PALE = RGBColor(248, 249, 251)
WHITE = RGBColor(255, 255, 255)
GOLD = RGBColor(185, 126, 0)


def set_fill(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency


def set_line(shape, color, width=1.0, transparency=0):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def add_text(slide, text, x, y, w, h, size=18, color=DARK, bold=False,
             align=PP_ALIGN.LEFT, font=FONT_CN, valign=MSO_ANCHOR.TOP,
             margin=0.03, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    p.line_spacing = 1.02
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def add_rich_lines(slide, items, x, y, w, h, size=16, color=DARK,
                   bullet=True, leading=1.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for idx, item in enumerate(items):
        if isinstance(item, tuple):
            txt, item_color, item_bold = item
        else:
            txt, item_color, item_bold = item, color, False
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = ("• " if bullet else "") + txt
        p.space_after = Pt(7)
        p.line_spacing = leading
        for run in p.runs:
            run.font.name = FONT_CN
            run.font.size = Pt(size)
            run.font.bold = item_bold
            run.font.color.rgb = item_color
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=LIGHT, radius=True, width=1.0):
    typ = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(s, fill)
    set_line(s, line, width)
    return s


def rule(slide, x, y, w, color=RED, h=0.035):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(s, color)
    s.line.fill.background()
    return s


def connector(slide, x1, y1, x2, y2, color=MID, width=1.5, dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    set_line(c, color, width)
    if dashed:
        c.line.dash_style = 4
    return c


def add_header(slide, title, page):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    slide.shapes.add_picture(str(LOGO), Inches(0.25), Inches(0.12), width=Inches(2.55))
    add_text(slide, title, 3.05, 0.18, 9.85, 0.58, size=22, bold=True,
             valign=MSO_ANCHOR.MIDDLE)
    rule(slide, 0.26, 0.91, 10.25, RED, 0.035)
    rule(slide, 10.62, 0.91, 2.45, TEAL, 0.035)
    add_text(slide, str(page), 12.75, 7.12, 0.25, 0.18, size=10, color=MID,
             align=PP_ALIGN.RIGHT)


def add_cover(slide, title, subtitle, paper, accent=TEAL, visual=None):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    slide.shapes.add_picture(str(LOGO), Inches(0.62), Inches(0.42), width=Inches(4.6))
    rule(slide, 0.0, 1.38, 10.55, RED, 0.12)
    rule(slide, 10.70, 1.38, 2.15, accent, 0.12)
    add_text(slide, title, 0.9, 2.02, 11.55, 1.1, size=31, color=RED,
             bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, subtitle, 1.3, 3.10, 10.7, 0.78, size=19, color=DARK,
             bold=False, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    if visual:
        add_picture_contain(slide, visual, 3.25, 3.88, 6.85, 1.75)
    else:
        # Compact authoritative/speculative timeline.
        rule(slide, 2.0, 4.42, 8.6, RED, 0.055)
        for i in range(6):
            node = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2.2+i*1.55), Inches(4.27), Inches(0.32), Inches(0.32))
            set_fill(node, RED if i < 4 else accent)
            node.line.fill.background()
        rule(slide, 3.85, 4.96, 5.25, accent, 0.055)
        add_text(slide, "Authoritative", 2.0, 4.60, 1.55, 0.3, size=12, color=RED, bold=True)
        add_text(slide, "Speculative", 3.85, 5.04, 1.4, 0.3, size=12, color=accent, bold=True)
    add_text(slide, paper, 0.75, 5.95, 11.8, 0.4, size=13, color=MID,
             align=PP_ALIGN.CENTER)
    add_text(slide, "汇报人：张帅    ｜    Agent 动作投机 Survey    ｜    2026.08", 0.75, 6.36, 11.8, 0.4,
             size=15, color=RED, bold=True, align=PP_ALIGN.CENTER)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.95), Inches(13.333), Inches(0.55))
    set_fill(band, RED)
    band.line.fill.background()
    add_text(slide, "自强 · 弘毅 · 求是 · 拓新", 0.5, 7.04, 12.3, 0.25,
             size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)


def add_card(slide, title, body, x, y, w, h, accent=TEAL, number=None, size=16):
    rect(slide, x, y, w, h, WHITE, LIGHT, True, 1.0)
    if number is not None:
        n = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+0.18), Inches(y+0.20), Inches(0.42), Inches(0.42))
        set_fill(n, accent)
        n.line.fill.background()
        add_text(slide, str(number), x+0.18, y+0.21, 0.42, 0.38, size=13, color=WHITE,
                 bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        tx = x + 0.72
        tw = w - 0.90
    else:
        tx = x + 0.18
        tw = w - 0.36
    add_text(slide, title, tx, y+0.16, tw, 0.34, size=size, color=accent, bold=True)
    add_text(slide, body, x+0.18, y+0.58, w-0.36, h-0.70, size=max(14, size-2), color=DARK)


def add_metric(slide, value, label, x, y, w=1.8, accent=TEAL):
    rect(slide, x, y, w, 0.86, PALE, accent, True, 1.0)
    add_text(slide, value, x+0.10, y+0.06, w-0.20, 0.42, size=23, color=accent,
             bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x+0.08, y+0.52, w-0.16, 0.23, size=11, color=MID,
             align=PP_ALIGN.CENTER)


def add_source(slide, text, x, y, w):
    add_text(slide, text, x, y, w, 0.20, size=9, color=MID, italic=True,
             align=PP_ALIGN.RIGHT)


def add_picture_contain(slide, path, x, y, w, h):
    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w/iw, h/ih)
    pw, ph = iw*scale, ih*scale
    px, py = x+(w-pw)/2, y+(h-ph)/2
    return slide.shapes.add_picture(str(path), Inches(px), Inches(py), width=Inches(pw), height=Inches(ph))


def raster_pdf(path, name, zoom=2.5):
    out = ASSET / f"{name}.png"
    if out.exists():
        return out
    doc = fitz.open(path)
    page = doc[0]
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=True)
    pix.save(out)
    doc.close()
    return out


def formula_png(name, formula, fontsize=26, width=8.5, height=0.8):
    out = ASSET / f"formula-{name}.png"
    if out.exists():
        return out
    matplotlib.rcParams["mathtext.fontset"] = "stix"
    matplotlib.rcParams["font.family"] = "Liberation Serif"
    fig = plt.figure(figsize=(width, height), dpi=220, facecolor=(1, 1, 1, 0))
    fig.text(0.5, 0.5, formula, ha="center", va="center", fontsize=fontsize, color="#1a2230")
    plt.axis("off")
    fig.savefig(out, transparent=True, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)
    return out


def set_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.clear()
    tf.paragraphs[0].text = text


def new_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def overview_slide(slide, title, bg_items, cards, metrics=None, footer=None, page=2):
    add_header(slide, title, page)
    add_text(slide, "背景 / 核心问题", 0.45, 1.18, 4.35, 0.36, size=18, color=RED, bold=True)
    rect(slide, 0.42, 1.58, 4.45, 4.72, PALE_RED, RED, True, 1.0)
    add_rich_lines(slide, bg_items, 0.67, 1.86, 3.98, 4.15, size=16)
    add_text(slide, "主要贡献", 5.17, 1.18, 3.0, 0.36, size=18, color=TEAL, bold=True)
    for i, (ct, cb) in enumerate(cards):
        add_card(slide, ct, cb, 5.15, 1.58+i*1.43, 7.65, 1.20, TEAL, i+1, 16)
    if metrics:
        total = len(metrics)
        mw = min(2.2, 7.2/total)
        start = 5.25
        for i, (value, label, accent) in enumerate(metrics):
            add_metric(slide, value, label, start+i*(mw+0.18), 5.92, mw, accent)
    if footer:
        rect(slide, 0.42, 6.52, 12.38, 0.42, WHITE, RED, True, 0.9)
        add_text(slide, footer, 0.62, 6.60, 11.98, 0.23, size=14, color=RED, bold=True, align=PP_ALIGN.CENTER)


def steps_slide(slide, title, steps, page, bottom=None, colors=None):
    add_header(slide, title, page)
    colors = colors or [RED, TEAL, BLUE, RED2, TEAL2]
    n = len(steps)
    y = 1.62
    for i, (st, sb) in enumerate(steps):
        x = 0.55 + i*(12.25/n)
        w = (12.25/n)-0.28
        add_card(slide, st, sb, x, y, w, 2.58, colors[i % len(colors)], i+1, 15)
        if i < n-1:
            connector(slide, x+w, y+1.29, x+w+0.25, y+1.29, colors[i % len(colors)], 2.0)
    if bottom:
        if isinstance(bottom, str):
            rect(slide, 0.55, 4.55, 12.1, 1.95, PALE, LIGHT, True, 1.0)
            add_text(slide, bottom, 0.85, 4.84, 11.5, 1.25, size=18, color=DARK, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        else:
            for i, (bt, bb, bc) in enumerate(bottom):
                add_card(slide, bt, bb, 0.60+i*4.10, 4.60, 3.85, 1.75, bc, None, 16)


def summary_slide(slide, title, takeaways, limits, position, page=8):
    add_header(slide, title, page)
    add_text(slide, "三句话带走", 0.52, 1.22, 2.5, 0.36, size=18, color=TEAL, bold=True)
    for i, item in enumerate(takeaways):
        add_card(slide, f"0{i+1}", item, 0.52+i*4.13, 1.62, 3.78, 1.55, TEAL if i != 1 else BLUE, None, 16)
    add_text(slide, "适用边界 / 限制", 0.52, 3.50, 3.0, 0.36, size=18, color=RED, bold=True)
    rect(slide, 0.50, 3.90, 7.55, 2.50, PALE_RED, RED, True, 1.0)
    add_rich_lines(slide, limits, 0.76, 4.17, 7.05, 2.0, size=15)
    add_text(slide, "Survey 技术链中的位置", 8.35, 3.50, 4.0, 0.36, size=18, color=BLUE, bold=True)
    rect(slide, 8.32, 3.90, 4.48, 2.50, PALE_BLUE, BLUE, True, 1.0)
    add_text(slide, position, 8.62, 4.18, 3.88, 1.90, size=17, color=DARK,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def figures_slide(slide, title, items, metrics, page=7, note=None):
    add_header(slide, title, page)
    n = len(items)
    x0, gap = 0.45, 0.24
    w = (12.43-gap*(n-1))/n
    for i, (path, caption) in enumerate(items):
        x = x0+i*(w+gap)
        rect(slide, x, 1.45, w, 4.35, WHITE, LIGHT, True, 1.0)
        add_picture_contain(slide, path, x+0.12, 1.62, w-0.24, 3.65)
        add_text(slide, caption, x+0.14, 5.28, w-0.28, 0.34, size=12, color=MID, align=PP_ALIGN.CENTER)
    for i, (v, lab, col) in enumerate(metrics):
        add_metric(slide, v, lab, 0.72+i*2.35, 6.02, 2.12, col)
    if note:
        add_text(slide, note, 8.0, 6.04, 4.55, 0.74, size=14, color=DARK,
                 valign=MSO_ANCHOR.MIDDLE)


# ---------------- DSP ----------------

def dsp_slide(prs, idx):
    s = new_slide(prs)
    if idx == 1:
        add_cover(s, "Dynamic Speculative Agent Planning",
                  "用在线预测的动态投机深度控制延迟—成本权衡",
                  "Yilin Guan et al. · arXiv:2509.01920")
        set_notes(s, "这篇工作不是重新设计动作 Speculator，而是回答一个更具体的系统问题：已经有快慢两个 Agent 并行规划时，每一轮到底应该投机多远？我会把重点放在动态 k、在线价值学习和延迟—成本控制。")
    elif idx == 2:
        overview_slide(s, "为什么固定投机步数 k 不够？",
            ["Agent 规划串行且延迟高；快 A 可草拟，慢 T 负责权威验证。",
             "k 太小：并行窗口没吃满；k 太大：失配后的草拟调用全部浪费。",
             "同一任务的不同阶段难度不同，不存在始终最优的固定 k。"],
            [("Pareto 分析", "将 k 作为延迟—成本的核心控制量。"),
             ("动态 k 预测", "DistilBERT 状态价值预测器在线学习连续命中长度。"),
             ("用户可控", "expectile 或 offset 将运行点在“更快—更省”间移动。")],
            footer="最终轨迹始终由 Target 决定；动态预测只改变提前执行多远。")
        set_notes(s, "固定 k 的问题是它用一个全局常数代替了局部可预测性。简单阶段可能连续命中很多步，复杂阶段第一步就失配。DSP 的三个贡献分别对应分析、预测和部署控制。")
    elif idx == 3:
        steps_slide(s, "固定 k：加速来自并行，浪费来自失配之后",
            [("A 草拟 k 步", "快 Approximation Agent 连续生成未来规划。"),
             ("T 并行验证", "慢 Target Agent 从同一状态生成权威步骤。"),
             ("首个 mismatch", "取消失配点之后所有未确认线程。"),
             ("跟随 T 恢复", "从 Target 的正确步骤继续，保持轨迹一致。")], 3,
            bottom=[("k < m", "冗余少，但留下未利用并行窗口。", BLUE),
                    ("k ≈ m", "接近实际连续命中长度，收益最好。", TEAL),
                    ("k > m", "延迟不再下降，token/API 冗余上升。", RED)])
        set_notes(s, "这一页先讲清固定 k 基线，否则后面动态预测没有参照。关键是首个失配后的草拟全部作废。因此最优 k 其实是对当前状态可连续命中长度 m 的估计。")
    elif idx == 4:
        add_header(s, "DSP：在每个投机 episode 中动态决定 k", 4)
        nodes = [("状态 s_i", 0.55, RED), ("Predictor\nVθ(s_i) → k_i", 2.55, BLUE),
                 ("Approximation A\n最多草拟 k_i 步", 5.05, TEAL),
                 ("Target T\n并行逐步验证", 7.95, RED),
                 ("Stop / Recover\n命中 k_i 或首个失配", 10.65, GOLD)]
        for i, (txt, x, col) in enumerate(nodes):
            rect(s, x, 1.75, 1.95, 1.25, WHITE, col, True, 1.4)
            add_text(s, txt, x+0.08, 1.96, 1.79, 0.82, size=15, color=col, bold=True,
                     align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
            if i < len(nodes)-1:
                connector(s, x+1.95, 2.37, nodes[i+1][1], 2.37, MID, 1.7)
        rect(s, 2.35, 3.55, 8.7, 1.35, PALE_BLUE, BLUE, True, 1.1)
        add_text(s, "Online Trainer（后台异步）", 2.62, 3.74, 2.45, 0.32, size=17, color=BLUE, bold=True)
        add_text(s, "用 Target 得到的实际连续命中长度构造 (s,k) → 更新 training copy → 热同步 inference copy", 5.0, 3.70, 5.75, 0.62, size=15, color=DARK)
        connector(s, 10.9, 3.0, 10.9, 4.22, BLUE, 1.3, True)
        connector(s, 2.55, 4.22, 2.55, 3.0, BLUE, 1.3, True)
        for i, (lab, col, txt) in enumerate([("A/T 主执行线", RED, "决定权威轨迹"), ("Predictor inference", TEAL, "与 A 当前步并行"), ("Trainer", BLUE, "不阻塞当前任务")]):
            add_card(s, lab, txt, 0.62+i*4.1, 5.38, 3.8, 1.15, col, None, 15)
        set_notes(s, "DSP 的运行时闭环分成三条线。A 和 T 负责实际规划；预测器在 episode 开始时估计 k，并且与 A 并行来隐藏自身延迟；训练器则等 T 给出标签后在后台更新。")
    elif idx == 5:
        add_header(s, "预测“还能连续命中几步”：状态价值建模", 5)
        f1 = formula_png("dsp-value", r"$V_{\pi}(s)=\mathbb{E}_{\pi}[G_t\mid s_t=s],\qquad G_t=\sum_{i=t}^{T-1}\gamma^{i-t}R(s_i,a_i)$", 25)
        f2 = formula_png("dsp-loss", r"$L_{\theta}=\mathbb{E}\left[(G_t^{\lambda}-V_{\theta}(s_t))^2\right]$", 30)
        f3 = formula_png("dsp-lambda", r"$G_t^{\lambda}=(1-\lambda)\sum_{n=1}^{T-t-1}\lambda^{n-1}G_{t:t+n}+\lambda^{T-t-1}G_t$", 24)
        add_card(s, "状态与回报", "s_t 是已验证轨迹；G_t 是到首次 mismatch 之前的实际连续命中步数。", 0.55, 1.45, 4.10, 1.35, RED, None, 16)
        rect(s, 4.95, 1.45, 7.85, 1.35, PALE, LIGHT, True, 1.0)
        add_picture_contain(s, f1, 5.20, 1.76, 7.35, 0.65)
        add_card(s, "TD / λ-return 目标", "DistilBERT 只回归 Vθ(s)，不负责生成 Agent 动作。", 0.55, 3.05, 4.10, 1.45, TEAL, None, 16)
        rect(s, 4.95, 3.05, 7.85, 1.45, PALE_TEAL, TEAL, True, 1.0)
        add_picture_contain(s, f2, 5.30, 3.25, 7.10, 0.55)
        add_picture_contain(s, f3, 5.18, 3.78, 7.35, 0.55)
        add_card(s, "λ = 0", "更接近 1-step TD，方差小、偏差大。", 0.65, 5.05, 3.7, 1.35, BLUE, None, 16)
        add_card(s, "0 < λ < 1", "在短期 bootstrap 与完整回报之间做权衡。", 4.82, 5.05, 3.7, 1.35, TEAL, None, 16)
        add_card(s, "λ = 1, γ = 1", "G_t = k，退化为在线监督回归可命中步数。", 8.97, 5.05, 3.7, 1.35, RED, None, 16)
        set_notes(s, "这里不要把 DSP 讲成一个复杂的策略学习问题。它实际上是状态价值回归：用当前轨迹预测还能命中多少步。当 lambda=1 且 gamma=1 时，就是直接用实际 k 做在线监督学习。")
    elif idx == 6:
        add_header(s, "一个预测器，两种“更激进 / 更保守”控制方式", 6)
        img = raster_pdf(DSP/"fig/openagi_react_cost_comp.pdf", "dsp-cost-control")
        add_card(s, "A｜Biased value / expectile", "用非对称损失让估计更偏高或偏低。\nτ > 0.5：更激进、更低延迟、更高成本。\nτ < 0.5：更保守、更少冗余。", 0.55, 1.50, 4.05, 2.25, RED, None, 16)
        add_card(s, "B｜Unbiased value + offset", "保持价值估计无偏，部署时用\nk = clip(round(Vθ(s)) + β)\nβ > 0 向速度端移动；β < 0 向成本端移动。", 0.55, 4.05, 4.05, 2.25, TEAL, None, 16)
        rect(s, 4.95, 1.50, 7.82, 4.82, WHITE, LIGHT, True, 1.0)
        add_picture_contain(s, img, 5.20, 1.72, 7.32, 4.18)
        add_source(s, "论文原图：OpenAGI ReAct cost comparison", 7.3, 5.92, 5.0)
        rect(s, 4.95, 6.35, 7.82, 0.42, PALE_BLUE, BLUE, True, 0.8)
        add_text(s, "每个 τ / β 都是 Pareto 平面上一个可配置运行点，不是准确率调参。", 5.15, 6.43, 7.42, 0.22, size=14, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        set_notes(s, "论文给出两种控制方式。expectile 是在训练目标上引入高估或低估偏好；offset 则是保持预测器不变，部署时直接加一个 beta。后者更像实际系统中的用户滑杆。")
    elif idx == 7:
        i1 = raster_pdf(DSP/"fig/openagi_react_gen_vs_time.pdf", "dsp-gen-time")
        i2 = raster_pdf(DSP/"fig/openagi_react_cost_comp.pdf", "dsp-cost-comp")
        figures_slide(s, "动态 k 追平最快 fixed-k，但减少无效调用",
                      [(i1, "生成量—延迟关系"), (i2, "成本—延迟 Pareto 对比")],
                      [("≈ fastest", "延迟", TEAL), ("-30%", "总成本（最多）", RED), ("-60%", "无效成本（最多）", BLUE)],
                      note="实验重点是成本和延迟；质量不降由“最终跟随 Target”保证。")
        set_notes(s, "实验在 OpenAGI 和 TravelPlanner 上比较串行、不同 fixed-k 和 DSP。主要结论是 DSP 保留了接近最快固定策略的延迟，但显著减少被取消的额外调用。这页不用反复讲 benchmark 准确率，因为无损性由执行语义得到。")
    elif idx == 8:
        summary_slide(s, "DSP 解决的是“投机多远”，不是“如何预测动作”",
            ["核心对象：动态投机深度 k，匹配状态的局部可预测性。",
             "方法关键：将连续命中长度写成状态价值，异步在线更新。",
             "系统意义：在不改变 Target 轨迹前提下，显式配置速度—成本点。"],
            ["未动态选择 Approximation Agent，仍受 A/T 速度比和一致性限制。",
             "严格步骤对齐可能拒绝语义等价规划。",
             "在线冷启动和跨域泛化证据仍有限。"],
            "Speculative Actions\n预测—验证—提交\n↓\nDSP\n在线动态深度 k")
        set_notes(s, "总结时要避免把 DSP 说成更好的动作预测器。它是在现有快慢 Agent 并行基础上计算投机深度。与 Speculative Actions 的关系是：前者定义通用执行语义，DSP 增加深度控制器。")
    return s


# ---------------- PASTE ----------------

def paste_slide(prs, idx):
    s = new_slide(prs)
    if idx == 1:
        add_cover(s, "PASTE", "Parallelizing Tool Execution and LLM Generation for Low-Latency Agent Serving",
                  "Yifan Sui et al. · arXiv:2603.18897")
        set_notes(s, "PASTE 是一篇非常适合从算法过渡到系统的论文。它不只讨论怎么猜下一工具，还讨论提前执行后如何隔离、调度，以及工具变快后怎么避免 GPU 突然拥塞。")
    elif idx == 2:
        overview_slide(s, "Agent Serving 的关键路径不只是 LLM，工具占 45%–57%",
            ["LLM 生成 → 工具调用 → 等待 → LLM 恢复，语义依赖被实现成物理串行。",
             "Deep research、coding、AI4Science 中，工具执行占 E2E 延迟 45%–57%。",
             "要隐藏 exposed tool time，必须在 LLM 尚未输出时得到具体 (tool, args)。"],
            [("Pattern Analyzer", "抽取重复控制流+隐式数据流，实例化未来调用。"),
             ("Tool Scheduler", "准入、隔离、确认/提升/丢弃，权威调用优先。"),
             ("LLM–Tool Co-Scheduler", "节奏化 returning sessions，不让工具侧收益变成 GPU 排队。")],
            metrics=[("-43.5%", "平均任务时间", RED), ("1.8×", "更低工具延迟", TEAL)],
            footer="预测命中 ≠ E2E 一定变快；必须联合考虑工具与 LLM 两个资源域。")
        set_notes(s, "这页用两个观察建立问题。第一，工具在关键路径上占比很高；第二，现有 Agent 虽然在语义上必须等上一步，但某些工具序列和参数来源是可预测的。PASTE 的贡献是把这个机会做成一个全栈系统。")
    elif idx == 3:
        img = raster_pdf(PASTE/"images/ATC/PASTE-v2.pdf", "paste-architecture")
        add_header(s, "PASTE 不是单一预测器，而是三个闭环调度组件", 3)
        rect(s, 0.45, 1.35, 8.15, 5.30, WHITE, LIGHT, True, 1.0)
        add_picture_contain(s, img, 0.62, 1.55, 7.82, 4.68)
        add_source(s, "论文原图：PASTE system overview", 3.5, 6.28, 4.7)
        steps = [("1  Mining", "历史 session 提取稳定工具转移和参数派生规则。", RED),
                 ("2  Instantiation", "用当前上下文填充 pattern，得到具体 (tool,args)。", TEAL),
                 ("3  Speculation", "选择性提前执行，结果先进隔离区。", BLUE),
                 ("4  Match", "权威调用到达时规范化比较 name+args。", RED),
                 ("5  Pacing", "使用工具侧信号控制 session 返回 LLM 的节奏。", TEAL)]
        for i, (t, b, c) in enumerate(steps):
            add_card(s, t, b, 8.88, 1.33+i*1.06, 3.95, 0.92, c, None, 14)
        rect(s, 8.88, 6.65, 3.95, 0.35, PALE_RED, RED, True, 0.8)
        add_text(s, "未确认结果对 Agent 不可见", 9.05, 6.71, 3.60, 0.20, size=13, color=RED, bold=True, align=PP_ALIGN.CENTER)
        set_notes(s, "这张总图的阅读顺序是从预测、到工具侧执行、再到 LLM 侧恢复。最重要的无损约束是：Pattern Analyzer 只产生候选，预执行结果先隔离，只有权威调用完整命中后才进入 session。")
    elif idx == 4:
        steps_slide(s, "从“经常这样做”到“现在可以提前执行”",
            [("控制流挖掘", "Search → Visit\nEdit → Test\nGrep → FileEditor"),
             ("数据流学习", "下一工具参数来自上一输出或当前 session 字段。"),
             ("在线实例化", "只当所需输入已出现，才生成具体 tool+args。"),
             ("调度器仲裁", "多个 pattern 冲突时，按收益、置信度和资源准入。")], 4,
            bottom="Search(\"DSP paper\") 已返回 URL → 从当前上下文抽取 URL → LLM 还在生成时预先执行 Visit(url)")
        set_notes(s, "Pattern Analyzer 与用一个小模型直接猜调用不同。它的预测来源是历史轨迹的重复结构，并且一定要把参数来源解析出来，否则只知道下一个是 Visit 也无法执行。")
    elif idx == 5:
        img = raster_pdf(PASTE/"images/ATC/speculative_tool_execution.pdf", "paste-tool-lifecycle")
        add_header(s, "提前运行不等于提前提交：完整的投机生命周期", 5)
        add_picture_contain(s, img, 0.45, 1.32, 7.40, 3.05)
        add_source(s, "论文原图：speculative tool execution", 3.5, 4.32, 4.0)
        stages = [("Candidate", TEAL), ("Admitted", TEAL), ("Running", BLUE), ("Reuse", TEAL), ("Promote", BLUE), ("Discard", RED)]
        for i, (lab, col) in enumerate(stages):
            x = 0.48 + i*1.22
            rect(s, x, 4.75, 1.02, 0.72, WHITE, col, True, 1.2)
            add_text(s, lab, x+0.03, 4.94, 0.96, 0.28, size=12, color=col, bold=True, align=PP_ALIGN.CENTER)
            if i < len(stages)-1:
                connector(s, x+1.02, 5.11, x+1.20, 5.11, MID, 1.2)
        add_card(s, "准入", "预计可隐藏工具时间 > 调度与资源机会成本。", 8.12, 1.42, 4.70, 1.17, TEAL, None, 15)
        add_card(s, "隔离", "结果进 speculative store，不写 authoritative session state。", 8.12, 2.80, 4.70, 1.17, BLUE, None, 15)
        add_card(s, "严格匹配", "工具名+规范化参数一致才能 reuse；同一运行 job 可 promote。", 8.12, 4.18, 4.70, 1.17, RED, None, 15)
        rect(s, 0.60, 5.82, 12.08, 0.67, PALE_RED, RED, True, 1.0)
        add_text(s, "安全：优先只读 / 幂等 / 可隔离工具；有副作用工具需显式策略、sandbox 或非变更版本。", 0.82, 6.02, 11.62, 0.26, size=14, color=RED, bold=True, align=PP_ALIGN.CENTER)
        set_notes(s, "这页要强调提前执行和提前提交是两件事。无损性依赖隔离，并且权威调用拥有资源优先级。Promote 是一个很实用的状态：真实调用到来时，同一投机 job 正在跑，不用重新发起。")
    elif idx == 6:
        img = raster_pdf(PASTE/"images/ATC/ablation.pdf", "paste-ablation")
        add_header(s, "工具侧变快后，为什么 E2E 反而可能变慢？", 6)
        chain = [("工具更早完成", TEAL), ("Returning sessions 突然增加", BLUE), ("GPU/KV 压力上升", RED), ("Queueing 抵消收益", RED)]
        for i, (lab, col) in enumerate(chain):
            x = 0.45+i*3.12
            rect(s, x, 1.38, 2.75, 0.82, WHITE, col, True, 1.2)
            add_text(s, lab, x+0.10, 1.60, 2.55, 0.34, size=15, color=col, bold=True, align=PP_ALIGN.CENTER)
            if i < 3:
                connector(s, x+2.75, 1.79, x+3.06, 1.79, MID, 1.5)
        rect(s, 0.45, 2.55, 6.15, 3.68, PALE_TEAL, TEAL, True, 1.0)
        add_text(s, "LLM–Tool Co-Scheduler", 0.75, 2.82, 5.55, 0.4, size=20, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
        add_rich_lines(s, ["观测工具完成、reuse/promotion 和已隐藏时间。",
                           "联合观测 active requests、decode batch、context/KV pressure。",
                           "GPU 过载时限制同时恢复数；利用不足时快速放行。",
                           "目标：将工具侧 overlap 真正转化为 task E2E 收益。"], 0.80, 3.35, 5.48, 2.50, size=15)
        rect(s, 6.88, 2.55, 5.92, 3.68, WHITE, LIGHT, True, 1.0)
        add_picture_contain(s, img, 7.08, 2.72, 5.52, 2.95)
        add_source(s, "论文消融：Tool-only / LLM-only / Full PASTE", 8.0, 5.74, 4.35)
        rect(s, 0.45, 6.48, 12.35, 0.42, PALE_BLUE, BLUE, True, 0.8)
        add_text(s, "Tool-only 不是完整解法；Agent Serving 的目标是跨 LLM 与 Tool 两个资源域的 task E2E。", 0.70, 6.56, 11.85, 0.23, size=14, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        set_notes(s, "这是 PASTE 最有系统味道的一页。工具更早返回会让大量会话更集中地进入 GPU，如果不做 pacing，可能工具时间减少但 LLM 排队上升。因此应用层预测必须与 serving 调度协同。")
    elif idx == 7:
        i1 = raster_pdf(PASTE/"images/ATC/E2E_overall.pdf", "paste-e2e")
        i2 = raster_pdf(PASTE/"images/ATC/ablation.pdf", "paste-ablation-2")
        figures_slide(s, "全栈协同才能把工具侧 overlap 转化为 E2E 收益",
                      [(i1, "E2E latency vs. serving baselines"), (i2, "Tool-only / LLM-only / Full PASTE")],
                      [("-43.5%", "平均 E2E（最多）", RED), ("-55.4%", "p99（最多）", BLUE), ("1.8×", "工具延迟改善", TEAL)],
                      note="Tool-only：LLM queue 65s / E2E 419s\nFull PASTE：45s / 270s")
        set_notes(s, "PASTE 在 deep research、coding 和 scientific-agent 工作负载上与 vLLM、Agentix、ORION 和 SpecFaaS 比较。最值得讲的不是单一平均数字，而是消融证明 Tool-only 会把瓶颈挪到 GPU 队列，只有完整 PASTE 同时改善两侧。")
    elif idx == 8:
        summary_slide(s, "PASTE 把“动作投机”推进为“Agent Serving 全栈并行化”",
            ["预测来源：重复轨迹中的控制流+数据流，不额外调用 LLM。",
             "无损核心：观测式预测、结果隔离、权威优先、严格 name+args。",
             "系统核心：工具侧提前执行必须与 LLM 侧 pacing 一起设计。"],
            ["依赖重复 Agent 轨迹和可从已知上下文提取的具体参数。",
             "低频工具、开放参数和长尾 pattern 收益有限。",
             "不可撤销的外部副作用工具不能直接投机执行。"],
            "Speculative Actions\n通用执行语义\n↓\nPASTE\n预测 + 工具执行 + 全栈调度")
        set_notes(s, "PASTE 在 Survey 中的价值是它告诉我们：动作投机不是只把一个小模型接在 Agent 旁边。真实部署需要参数实例化、工具隔离和跨资源域调度。它的主要边界则是对重复 pattern 和可安全提前执行工具的依赖。")
    return s


# ---------------- SPORK ----------------

def spork_slide(prs, idx):
    s = new_slide(prs)
    fig = SPORK/"figures"
    if idx == 1:
        add_cover(s, "SPORK", "Self-Speculative Forking to Accelerate Agentic LLM Inference",
                  "Huajun Bai et al. · arXiv:2607.03333")
        set_notes(s, "SPORK 的切入点是：不用辅助模型、不依赖历史轨迹，而是让正在运行的 Actor 从自己的 Prefix KV 上分叉，强制探测下一个工具调用。这也是动作投机与 token speculative decoding 的桥接。")
    elif idx == 2:
        overview_slide(s, "外部 Speculator 的成本，也是 Agent 投机的系统瓶颈",
            ["外部小模型需要第二份权重/KV cache，且常与 Actor 的真实调用分布不一致。",
             "历史 pattern 对长尾工具和少见参数容易失效。",
             "研究问题：运行中的 LLM 能否足够早、足够准且失败后可回收地预测自己？"],
            [("D1｜Prefix-KV Fork", "首 token 后强制工具前缀，缓存命中将 probe 约 1.6s 降到 0.35s。"),
             ("D2｜Confidence Gate", "用 probe 自身 tool-name logprob，选择最早可信提交时刻。"),
             ("D3｜Token Reuse", "严格拒绝时，把共享工具调用前缀当作 token 草稿。")],
            footer="Training-free self-speculation：同一 Actor 、共享 Prefix KV、工具级严格匹配。")
        set_notes(s, "SPORK 将三个问题分开处理：D1 降低多跑一个 probe 的开销，D2 决定什么时候真正发工具，D3 让失配分支也能回收一部分 token。三者分别对应成本模型的不同项。")
    elif idx == 3:
        add_header(s, "同一模型确实“早知道、知道何时可信、失败也不全浪费”", 3)
        imgs = [fig/"fig1b_per_turn_fork_accuracy_top.png", fig/"fig_insight2_cot_curve_gaia.png", fig/"fig_insight3_prefix_distribution.png"]
        caps = [("Insight 1｜早知道工具名", "tool-name 74.6%–99.6%\nargs exact 7.6%–83.9%", RED),
                ("Insight 2｜信心随 CoT 上升", "GAIA args exact\n7.6% → 97.5%", TEAL),
                ("Insight 3｜拒绝仍有前缀", "中位 18 tokens\nP25 = 14 tokens", BLUE)]
        for i, path in enumerate(imgs):
            x = 0.40+i*4.30
            rect(s, x, 1.35, 4.08, 5.18, WHITE, LIGHT, True, 1.0)
            add_text(s, caps[i][0], x+0.18, 1.55, 3.72, 0.38, size=17, color=caps[i][2], bold=True, align=PP_ALIGN.CENTER)
            add_picture_contain(s, path, x+0.14, 2.02, 3.80, 3.55)
            add_text(s, caps[i][1], x+0.25, 5.70, 3.58, 0.58, size=16, color=caps[i][2], bold=True, align=PP_ALIGN.CENTER)
        set_notes(s, "三个经验发现支撑三个设计。模型在很早时就知道工具名，但参数需要更多 CoT 才稳定，所以需要动态置信门控。同时，即使完整调用不匹配，两者通常仍共享一段较长 token 前缀。")
    elif idx == 4:
        add_header(s, "SPORK 的“分叉—门控—严格验收—前缀回收”", 4)
        path = fig/"fig2_method_overview_v2.png"
        rect(s, 0.36, 1.25, 12.58, 5.72, WHITE, LIGHT, True, 1.0)
        add_picture_contain(s, path, 0.52, 1.42, 12.26, 5.20)
        add_source(s, "论文原图：SPORK method overview", 8.0, 6.66, 4.45)
        set_notes(s, "阅读这张图时先看 baseline 的串行工具等待，再看 SPORK 在首 token 之后分叉。D2 只在置信度足够高时发工具；主流到达后完整匹配 name+args，命中就用已完成结果，失配则串行执行并用 D3 回收 token 前缀。")
    elif idx == 5:
        add_header(s, "D1 先把“多跑一次模型”的前缀开销压到足够低", 5)
        f1 = formula_png("spork-cost", r"$\mathrm{Speedup}=\frac{T_{base}}{T_{base}^{*}-\alpha t_{overlap}+T_{oh}}$", 30)
        f2 = formula_png("spork-break", r"$\alpha\,t_{overlap}\;\geq\;T_{oh}$", 34, 4.2, 0.7)
        rect(s, 0.48, 1.32, 7.10, 1.18, PALE, LIGHT, True, 1.0)
        add_picture_contain(s, f1, 0.75, 1.55, 6.55, 0.56)
        rect(s, 7.82, 1.32, 5.00, 1.18, PALE_TEAL, TEAL, True, 1.0)
        add_text(s, "Break-even", 8.05, 1.48, 1.25, 0.30, size=16, color=TEAL, bold=True)
        add_picture_contain(s, f2, 9.10, 1.48, 3.30, 0.50)
        add_card(s, "α", "严格命中并接受已执行工具的 turn 比率。", 0.50, 2.80, 3.70, 1.30, RED, None, 16)
        add_card(s, "t_overlap", "工具执行与主生成真正重叠的时间。", 4.42, 2.80, 3.70, 1.30, TEAL, None, 16)
        add_card(s, "T_oh", "probe prefill/decode 以及对主流的干扰。", 8.34, 2.80, 3.70, 1.30, BLUE, None, 16)
        img = fig/"fig_d1_overhead.png"
        rect(s, 0.50, 4.42, 7.65, 2.18, WHITE, LIGHT, True, 1.0)
        add_picture_contain(s, img, 0.70, 4.55, 7.25, 1.75)
        add_card(s, "D1 的实际效果", "等 main 第一个 token 后再发 probe，前缀 KV 已入缓存。\nprefill 1.3s → 0.05s\n整个 probe 1.6s → 0.35s\nmain TPOT 干扰约 +0.22%", 8.42, 4.42, 4.38, 2.18, TEAL, None, 16)
        set_notes(s, "这个成本模型是理解 SPORK 收益的关键。只有命中概率乘真实重叠时间大于 probe 开销，投机才值得。D1 通过首 token 后分叉复用 Prefix KV，把最重的重复 prefill 消掉。")
    elif idx == 6:
        add_header(s, "D2 提高期望重叠，D3 降低失败分支的浪费", 6)
        f = formula_png("spork-confidence", r"$c=\min_{i=2,\ldots,L}\exp(\ell_i)$", 31, 5.0, 0.7)
        rect(s, 0.45, 1.35, 6.05, 5.40, PALE_TEAL, TEAL, True, 1.1)
        add_text(s, "D2｜Adaptive Confidence Gate", 0.72, 1.60, 5.50, 0.42, size=20, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
        add_picture_contain(s, f, 1.10, 2.08, 4.70, 0.64)
        add_picture_contain(s, fig/"fig_d2_confidence_gate.png", 0.75, 2.82, 5.45, 2.70)
        add_text(s, "θ=0.90：Precision 88%，Recall 100%，F1=0.937；过滤 77% probes", 0.72, 5.72, 5.50, 0.50, size=14, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
        rect(s, 6.80, 1.35, 6.05, 5.40, PALE_BLUE, BLUE, True, 1.1)
        add_text(s, "D3｜Partial-token Accept", 7.08, 1.60, 5.50, 0.42, size=20, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        d3 = raster_pdf(fig/"fig_d3_spec_dec.pdf", "spork-d3")
        add_picture_contain(s, d3, 7.15, 2.10, 5.35, 2.65)
        add_rich_lines(s, ["主流进入 tool-call opener 时，将 probe body 注入为草稿。",
                           "主模型并行验证，接受最长一致前缀 k。",
                           "工具级失配仍获得 token 级复用。"], 7.15, 4.95, 5.25, 1.50, size=14)
        set_notes(s, "D2 使用工具名 span 中最低的 top-1 token 概率作为保守置信度，低于阈值时并不真正发工具。D3 则不放宽工具级严格匹配，而是把失配 probe 当成 token 草稿，因此不会改变工具语义。")
    elif idx == 7:
        figures_slide(s, "SPORK 的收益不大但稳健：主要改善真实工具任务 P95",
                      [(fig/"fig_hero_e2e_speedup.png", "真实 Agent 任务 E2E/P95"), (fig/"fig_operating_regime_map.png", "何时 self-speculation 划算")],
                      [("-6%∼18%", "P95 延迟", TEAL), ("±1 pp", "任务准确率", BLUE), ("5", "主要 benchmarks", RED)],
                      note="有利：工具慢 + CoT 窗口长 + 严格命中高\n不利：工具快 + 参数晚确定")
        set_notes(s, "SPORK 的收益比 PASTE 小，但它的优点是训练免费、不额外常驻模型，且对主轨迹使用严格匹配。收益最大的情况是工具较慢、剩余 CoT 足够长，并且完整参数可以较早确定。")
    elif idx == 8:
        summary_slide(s, "SPORK 是 Agent 动作投机与 Token 投机解码之间最直接的桥",
            ["预测源：当前 Actor 自身 Prefix-KV fork，无额外模型与数据。",
             "两层接受：工具级完整 name+args；失败分支回到 token 级验证。",
             "系统经验：用 α·t_overlap 与 T_oh 的 break-even 判断是否划算。"],
            ["工具名早确定不等于完整参数早确定，严格命中仍是瓶颈。",
             "D3 需修改推理引擎 proposer path；HTTP 控制器只能实现 D1+D2。",
             "对副作用工具，置信门控不能替代隔离/确认。"],
            "PASTE\n历史轨迹模式\n↓\nSPORK\n同模型 fork + token 前缀回收")
        set_notes(s, "SPORK 的关键意义是把同一个 probe 在两个粒度上复用：完整结构命中时复用工具结果，不命中时复用 token 前缀。它也明确告诉我们，自投机不是总能赚，必须满足重叠收益大于 probe 开销。")
    return s


# ---------------- Joint RL ----------------

def joint_slide(prs, idx):
    s = new_slide(prs)
    fig = JOINT/"figures"
    if idx == 1:
        add_cover(s, "Speculate While You Reason", "Teaching Agents to Predict Their Next Tool Call via Joint Agent–Speculator RL",
                  "arXiv:2607.25816", visual=fig/"intro_self_speculation.png")
        set_notes(s, "这篇论文与 SPORK 使用同一模型的出发点相同，但路线相反。SPORK 不训练，只探测当前 Actor；这篇论文则把预测自己的下一个工具调用直接变成联合训练目标。")
    elif idx == 2:
        overview_slide(s, "小外部 Speculator 不一定既准又便宜，关键是与 Agent 对齐",
            ["外部小模型带来第二份权重/KV，但仍常与 4B Agent 真实调用分布不一致。",
             "通用 Agent 未被教过“从部分轨迹猜自己的下一调用”。",
             "Agent policy 会持续更新，预测目标也会漂移，离线旧轨迹很快过时。"],
            [("Self-speculating agent", "同一 πθ 以固定后缀切换 Agent / Speculator 模式，共享权重与 Prefix KV。"),
             ("On-policy targets", "从当前 Agent rollout 截取部分轨迹，实际下一调用作为目标。"),
             ("Joint RL recipe", "双模式 SFT + 4:8 交替 DAPO + 切换模式时重置 optimizer。")],
            metrics=[("44.1→61.2", "Qwen3-4B Hit@1", RED), ("48.9→66.3", "Qwen3.5-4B Hit@1", TEAL)],
            footer="匹配训练域的 task success 保持；但训练修改了 Agent policy，不是严格轨迹无损。")
        set_notes(s, "这页将问题分成成本、分布对齐和目标漂移三个层面。论文的解法是同一模型双模式，并且从当前 policy 的新鲜 rollout 中构造投机目标。这也意味着它不能简单归类为严格无损系统。")
    elif idx == 3:
        add_header(s, "“更小”的外部模型并不自然意味“更好”的 Speculator", 3)
        path = fig/"intro_self_speculation.png"
        rect(s, 0.40, 1.27, 8.25, 5.58, WHITE, LIGHT, True, 1.0)
        add_picture_contain(s, path, 0.54, 1.43, 7.97, 5.14)
        add_source(s, "论文原图：external draft vs. self-speculation", 3.9, 6.59, 4.4)
        add_card(s, "External small model", "额外权重 + 独立 Prefix KV\n预测分布与 Target 不一致\nQwen3 最强外部小模型平均 Hit@1 约 18.6", 8.92, 1.38, 3.90, 2.10, RED, None, 16)
        add_card(s, "Same-model self-speculation", "共享权重 + Prefix KV\n与自身动作分布更接近\nQwen3-4B 未训练自投机平均 Hit@1 约 29.1", 8.92, 3.78, 3.90, 2.10, TEAL, None, 16)
        rect(s, 8.92, 6.15, 3.90, 0.56, PALE_BLUE, BLUE, True, 0.8)
        add_text(s, "需学习：h_t → 下一结构调用 a_t", 9.10, 6.32, 3.54, 0.22, size=14, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        set_notes(s, "外部小模型任务能力强并不等于它能模仿目标 Agent 的下一调用。论文的实验表明，未训练的同模型自投机已经比外部小模型更好。因此下一步是专门训练这个部分轨迹到下一调用的映射。")
    elif idx == 4:
        add_header(s, "同一条 Agent rollout 同时提供任务学习和投机目标", 4)
        path = fig/"example-speculation.png"
        rect(s, 0.42, 1.25, 7.30, 5.62, WHITE, LIGHT, True, 1.0)
        add_picture_contain(s, path, 0.56, 1.42, 7.02, 5.16)
        add_source(s, "论文原图：partial trajectory 与 predicted call", 3.2, 6.60, 4.2)
        steps = [("1", "用当前 πθ 生成完整 Agent rollout τ。", RED),
                 ("2", "对每个工具轮 t 截取调用前缀 h_t。", BLUE),
                 ("3", "追加固定投机后缀 s_sp，触发 Speculator mode。", TEAL),
                 ("4", "目标就是同一 rollout 中 policy 实际产生的 a_t。", TEAL)]
        for i, (num, txt, col) in enumerate(steps):
            add_card(s, f"Step {num}", txt, 8.03, 1.28+i*1.32, 4.78, 1.08, col, None, 15)
        rect(s, 8.03, 6.60, 4.78, 0.38, PALE_BLUE, BLUE, True, 0.8)
        add_text(s, "Agent：长 reasoning trajectory｜Speculator：单个 structured call", 8.20, 6.68, 4.44, 0.20, size=12, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        set_notes(s, "数据构造的核心是 on-policy。每次 Agent 产生新 rollout 后，同一批数据一方面用最终答案奖励训 Agent，另一方面从每个工具轮截取部分轨迹，让模型直接预测它自己真正会发出的下一调用。")
    elif idx == 5:
        add_header(s, "部署需要完全匹配，训练则需要可学的稠密信号", 5)
        f1 = formula_png("joint-name", r"$R_{name}=\mathbf{1}[\hat n_t=n_t]$", 31, 4.8, 0.7)
        f2 = formula_png("joint-args", r"$R_{args}=\frac{1}{|\mathrm{keys}(\sigma_t)|}\sum_{k\in\mathrm{keys}(\sigma_t)}\mathrm{F1}(\hat\sigma_t[k],\sigma_t[k])$", 23)
        f3 = formula_png("joint-reward", r"$R_{sp}=R_{name}\cdot R_{args}$", 34, 4.8, 0.7)
        add_card(s, "Tool name", "", 0.52, 1.40, 3.75, 1.58, RED, None, 16)
        add_picture_contain(s, f1, 0.75, 1.92, 3.30, 0.42)
        add_text(s, "工具名错误立即得 0，与“结果无法复用”一致。", 0.74, 2.45, 3.32, 0.34, size=13, color=DARK, align=PP_ALIGN.CENTER)
        add_card(s, "Arguments", "", 4.52, 1.40, 4.30, 1.58, TEAL, None, 16)
        add_picture_contain(s, f2, 4.72, 1.91, 3.90, 0.48)
        add_text(s, "名字命中后按 key 做 macro token-F1，近似值获得局部信号。", 4.76, 2.45, 3.82, 0.34, size=13, color=DARK, align=PP_ALIGN.CENTER)
        add_card(s, "Final reward", "", 9.05, 1.40, 3.75, 1.58, BLUE, None, 16)
        add_picture_contain(s, f3, 9.28, 1.92, 3.28, 0.42)
        add_text(s, "只有工具名正确，参数稠密奖励才生效。", 9.28, 2.45, 3.28, 0.34, size=13, color=DARK, align=PP_ALIGN.CENTER)
        examples = [("工具名错", "search → visit", "R_sp = 0", RED),
                    ("工具名对，参数近似", "query 字符串部分重叠", "0 < R_sp < 1", TEAL),
                    ("完全命中", "name + argument dictionary exact", "Hit@1 = 1", BLUE)]
        for i, (t, b, v, c) in enumerate(examples):
            add_card(s, t, f"{b}\n{v}", 0.58+i*4.17, 3.22, 3.88, 1.65, c, None, 15)
        rect(s, 0.58, 5.26, 12.10, 1.18, PALE_RED, RED, True, 1.0)
        add_text(s, "训练标准", 0.86, 5.50, 1.55, 0.34, size=17, color=RED, bold=True)
        add_text(s, "参数 token-F1 提供 partial credit，让 RL 有足够稠密的信号。", 2.30, 5.49, 4.40, 0.38, size=15, color=DARK)
        add_text(s, "部署复用标准", 7.05, 5.50, 1.90, 0.34, size=17, color=BLUE, bold=True)
        add_text(s, "解析后 tool name + 完整 argument dictionary 精确一致。", 8.90, 5.49, 3.45, 0.58, size=15, color=DARK)
        set_notes(s, "训练时如果只用完全匹配 0/1 奖励，信号会过于稀疏。所以论文在工具名正确后，对每个参数值给 token-F1 partial credit。但要特别说明，部署复用结果时仍要完整精确匹配，稠密奖励没有放宽验收标准。")
    elif idx == 6:
        add_header(s, "数据收集耦合，优化步骤分离：一个 policy 交替学两个目标", 6)
        # Circular 4:8 training loop.
        rect(s, 0.55, 1.40, 7.35, 5.30, WHITE, LIGHT, True, 1.0)
        add_text(s, "Joint Agent–Speculator RL", 0.90, 1.68, 6.65, 0.42, size=21, color=DARK, bold=True, align=PP_ALIGN.CENTER)
        add_card(s, "SFT warmup", "成功 Agent 轨迹 + 部分轨迹投机样本", 2.40, 2.22, 3.65, 1.05, BLUE, None, 15)
        add_card(s, "4 × Agent updates", "R_ag：最终任务奖励\nDAPO grouped advantages", 0.95, 3.75, 2.65, 1.42, RED, None, 16)
        add_card(s, "8 × Spec. updates", "R_sp：下一工具奖励\nDAPO grouped advantages", 4.85, 3.75, 2.65, 1.42, TEAL, None, 16)
        connector(s, 3.60, 4.46, 4.85, 4.46, BLUE, 2.0)
        connector(s, 6.18, 3.75, 6.18, 3.28, TEAL, 1.5)
        connector(s, 2.28, 3.28, 2.28, 3.75, RED, 1.5)
        rect(s, 2.52, 5.48, 3.65, 0.76, PALE_BLUE, BLUE, True, 1.0)
        add_text(s, "每次模式切换重置 optimizer state", 2.72, 5.70, 3.25, 0.28, size=14, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        add_card(s, "稳定技巧 1", "先 SFT 建立两种正确输出格式。", 8.25, 1.42, 4.55, 1.28, BLUE, 1, 16)
        add_card(s, "稳定技巧 2", "切换时清除 momentum / adaptive statistics 干扰。", 8.25, 2.98, 4.55, 1.28, TEAL, 2, 16)
        add_card(s, "稳定技巧 3", "Agent reward 惩罚无意义重复工具调用。", 8.25, 4.54, 4.55, 1.28, RED, 3, 16)
        rect(s, 8.25, 6.10, 4.55, 0.58, PALE_RED, RED, True, 0.8)
        add_text(s, "不将 R_ag 和 R_sp 粗暴线性相加后同步更新", 8.47, 6.25, 4.11, 0.26, size=13, color=RED, bold=True, align=PP_ALIGN.CENTER)
        set_notes(s, "论文把数据收集和优化策略分开。同一批 rollout 同时提供两种数据，但每个更新 block 只优化一个目标。最终采用 4 次 Agent 更新加 8 次 Speculator 更新的循环，并在模式切换时重置优化器状态。")
    elif idx == 7:
        abl = raster_pdf(fig/"reward_ablation_trends.pdf", "joint-ablation")
        add_header(s, "联合 RL 把自投机从“可用起点”提高到“更高严格命中率”", 7)
        rect(s, 0.45, 1.32, 6.12, 4.65, WHITE, LIGHT, True, 1.0)
        add_text(s, "Base → SFT → Joint RL", 0.75, 1.60, 5.52, 0.40, size=20, color=DARK, bold=True, align=PP_ALIGN.CENTER)
        data = [("Qwen3-4B", [29.1, 44.1, 61.2], [26.4, 26.6, 27.7], RED),
                ("Qwen3.5-4B", [33.2, 48.9, 66.3], [49.0, 49.2, 50.6], TEAL)]
        for row, (name, vals, succ, col) in enumerate(data):
            y = 2.28 + row*1.60
            add_text(s, name, 0.78, y+0.20, 1.35, 0.34, size=15, color=col, bold=True)
            labs = ["Base", "SFT", "RL"]
            for i, v in enumerate(vals):
                x = 2.15+i*1.28
                rect(s, x, y, 1.05, 0.82, PALE, col, True, 1.0)
                add_text(s, f"{v:.1f}", x+0.05, y+0.10, 0.95, 0.35, size=20, color=col, bold=True, align=PP_ALIGN.CENTER)
                add_text(s, labs[i], x+0.05, y+0.52, 0.95, 0.20, size=11, color=MID, align=PP_ALIGN.CENTER)
                if i < 2:
                    connector(s, x+1.05, y+0.41, x+1.26, y+0.41, col, 1.4)
            add_text(s, f"Task success: {succ[0]:.1f} → {succ[1]:.1f} → {succ[2]:.1f}", 2.13, y+0.95, 3.65, 0.26, size=13, color=MID, align=PP_ALIGN.CENTER)
        rect(s, 6.85, 1.32, 5.98, 4.65, WHITE, LIGHT, True, 1.0)
        add_picture_contain(s, abl, 7.02, 1.58, 5.64, 3.95)
        add_source(s, "论文原图：reward ablation trends", 8.7, 5.62, 3.6)
        rect(s, 0.45, 6.25, 12.38, 0.58, PALE_BLUE, BLUE, True, 0.8)
        add_text(s, "移除 SFT warmup、optimizer reset 或 alternating schedule，都会使 Agent / Speculator reward 更弱或更不稳定。", 0.75, 6.41, 11.78, 0.27, size=14, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        set_notes(s, "主表要讲两个趋势：第一，SFT 已经明显提高自投机 Hit@1；第二，交替 RL 在五个 benchmark 上继续同向提高，并且匹配域任务成功率没有下降。右侧消融则支持三个稳定技巧都必要。")
    elif idx == 8:
        summary_slide(s, "该方法优化 Speculator–Agent 对齐，代价是改变 Agent 本身",
            ["架构：同一模型双模式，消除额外权重并共享 Prefix KV。",
             "学习：下一工具目标来自当前 Agent 新鲜 rollout，交替 RL 维持两种行为。",
             "结果：严格 Hit@1 大幅上升，匹配训练域的 task success 保持。"],
            ["不是 Speculative Actions 意义下的严格轨迹无损：RL 直接改了 Agent policy。",
             "跨域 Hit@1 可能上升，但任务成功率不总是稳定，仍敏感于训练分布。",
             "主要证据是 Hit@1 与 Agent 质量；全栈 E2E 延迟/资源证据不如 PASTE、SPORK 完整。"],
            "SPORK\n训练免费的同模型 probe\n↓\nJoint RL\n主动训练 Agent 学会预测自己")
        set_notes(s, "最后一定要把无损边界说清。训练后的 Agent 已不是原始串行 Agent，所以无法声称轨迹完全一致。它能声称的是匹配域实验上任务成功率保持。因此这篇论文在 Survey 中应放在“训练式自投机”而不是“严格无损执行”。")
    return s


BUILDERS = {"DSP": dsp_slide, "PASTE": paste_slide, "SPORK": spork_slide, "Joint-RL": joint_slide}
FILENAMES = {
    "DSP": "26.8-Agent动作投机-DSP-张帅.pptx",
    "PASTE": "26.8-Agent动作投机-PASTE-张帅.pptx",
    "SPORK": "26.8-Agent动作投机-SPORK-张帅.pptx",
    "Joint-RL": "26.8-Agent动作投机-Joint-RL-张帅.pptx",
}


def build(paper, phase):
    ASSET.mkdir(parents=True, exist_ok=True)
    DECK.mkdir(parents=True, exist_ok=True)
    part = DECK / (FILENAMES[paper].replace(".pptx", "-P1-5.pptx"))
    final = DECK / FILENAMES[paper]
    if phase == 1:
        prs = new_presentation()
        for idx in range(1, 6):
            BUILDERS[paper](prs, idx)
        prs.core_properties.title = f"{paper} - Agent 动作投机组会汇报"
        prs.core_properties.author = "张帅"
        prs.save(part)
        return part
    if not part.exists():
        raise FileNotFoundError(f"请先运行 phase 1: {part}")
    prs = Presentation(part)
    for idx in range(6, 9):
        BUILDERS[paper](prs, idx)
    prs.save(final)
    return final


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--paper", choices=list(BUILDERS), required=True)
    parser.add_argument("--phase", type=int, choices=[1, 2], required=True)
    args = parser.parse_args()
    print(build(args.paper, args.phase))


if __name__ == "__main__":
    main()
